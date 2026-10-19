"""Build REPORT.docx from REPORT_GDOC.md structure with local PNGs (upload to Drive → Open as Google Doc)."""

from pathlib import Path

from docx import Document
from docx.shared import Inches

DIR = Path(__file__).resolve().parent

sections = [
    (
        "CRIU capture latency: where the time goes",
        None,
        "Notes from local runs on one GPU node with vLLM (Qwen3). Goal: see where capture "
        "seconds go and what might be worth optimizing, not pick a single best storage setup.\n\n"
        "Issue: Reduce scale-dependent CRIU capture latency.",
    ),
    (
        "At a glance (14B model)",
        None,
        None,
    ),
]

TABLE_GLANCE = [
    ("Typical capture on the fast baseline path", "~19.5 s (median of 3 runs)"),
    ("Time writing the checkpoint image", "~11 s"),
    ("Time after the image is finished", "~8 s (unchanged across disk settings)"),
    ("Tail vs checkpoint size", "~0.15 s per GB (0.6B–14B)"),
    ("Leave engine running (diagnostic)", "Tail 8 s → 0 s; capture 20 s → 11 s"),
    ("Newer agent, staged copy", "~58 s total (~19 s copy)"),
    ("Newer agent, direct publish", "~24 s total"),
]

IMAGES = [
    ("Experiments in order", "0-experiment-roadmap.png"),
    ("What one capture looks like", "0-capture-timeline.png"),
    ("Trial 1: 14B capture by path", "1-phases-14b.png"),
    ("Tail vs model size", "2-tail-vs-size.png"),
    ("Trial 2: tail probe", "3-tail-probe.png"),
    ("Leave running vs control", "4-leaverunning.png"),
]


def main():
    doc = Document()
    doc.add_heading(sections[0][0], 0)
    doc.add_paragraph(sections[0][2])

    doc.add_heading("At a glance (14B model)", level=1)
    t = doc.add_table(rows=1 + len(TABLE_GLANCE), cols=2)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text = "What we measured"
    hdr[1].text = "Result"
    for i, (a, b) in enumerate(TABLE_GLANCE, start=1):
        t.rows[i].cells[0].text = a
        t.rows[i].cells[1].text = b

    doc.add_paragraph(
        "Restore was ~36 s on baseline and ~25–27 s on the newer stack (scale comparison only)."
    )

    for title, img in IMAGES:
        doc.add_heading(title, level=1)
        path = DIR / img
        if path.exists():
            doc.add_picture(str(path), width=Inches(6.2))
        else:
            doc.add_paragraph(f"(missing {img})")

    doc.add_heading("Trial 1 matrix (14B)", level=1)
    rows = [
        ("Direct I/O, 96 Gi", "19.5 s", "10.8 s", "8.7 s", "36 s"),
        ("Buffered, 4 Gi", "31.8 s", "22.8 s", "8.8 s", "—"),
        ("PageBroker + direct", "35.4 s", "26.6 s", "8.8 s", "64 s"),
    ]
    t2 = doc.add_table(rows=1 + len(rows), cols=5)
    t2.style = "Table Grid"
    for j, h in enumerate(["Path", "Capture", "Write", "Tail", "Restore"]):
        t2.rows[0].cells[j].text = h
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            t2.rows[i].cells[j].text = val

    doc.add_heading("Trial 3 (newer stack, 14B, n=1)", level=1)
    t3 = doc.add_table(rows=4, cols=5)
    t3.style = "Table Grid"
    for j, h in enumerate(["Step", "CRIU dump", "Publish", "Capture total", "Restore"]):
        t3.rows[0].cells[j].text = h
    data = [
        ("Baseline median", "19.5 s", "rename", "~19.5 s", "36 s"),
        ("3a staged publish", "34.8 s", "18.8 s copy", "57.7 s", "27 s"),
        ("3b direct publish", "21.2 s", "~0 s", "24.2 s", "25 s"),
    ]
    for i, row in enumerate(data, 1):
        for j, val in enumerate(row):
            t3.rows[i].cells[j].text = val

    doc.add_heading("What looks worth optimizing", level=1)
    for item in [
        "Post-write reap (~8 s at 14B) on the baseline path",
        "Staged publish copy (~19–20 s) before direct publish",
        "Buffered I/O with a 4 Gi agent",
        "CUDA checkpoint (~3 s), secondary",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    out = DIR / "REPORT.docx"
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
